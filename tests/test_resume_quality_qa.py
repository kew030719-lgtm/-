from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from career_radar.resume_qa import inspect_export
from career_radar.resume import build_profile
from career_radar.schemas import (
    CandidateContact,
    CandidateProfile,
    Evidence,
    JobSnapshot,
    ResumeBullet,
    ResumeDraftVersion,
    ResumeEntry,
    ResumeSection,
    ResumeSourceEntry,
    TargetJob,
)
from career_radar.tailoring import _public_profile, _repair_draft_source_metadata, supplement_entry_metadata


def _version(target: TargetJob) -> ResumeDraftVersion:
    return ResumeDraftVersion(
        version_id="version_1", draft_id="draft_1", version=1, tailoring_id="tailor_1",
        profile_id=target.profile_id, target_job_id=target.target_job_id,
        headline=target.title, contact=CandidateContact(
            profile_id=target.profile_id, name="张三", phone="13800138000",
            email="zhangsan@example.com",
        ), created_at=datetime.now(UTC).isoformat(),
    )


def test_confirmed_project_answer_gets_a_project_heading():
    answer = (
        "使用 Python 编写股票信息数据采集程序，负责数据请求、解析、清洗与保存，"
        "项目代码仓库 https://github.com/kew030719-lgtm/stock"
    )
    assert supplement_entry_metadata(answer) == ("project", "股票信息数据采集程序")

    profile = CandidateProfile(
        profile_id="profile_1",
        evidence=[Evidence(
            source_type="resume", source_id="profile_1", block_id="supplement_1",
            quote=answer, section="用户补充", provenance="user_confirmed",
        )],
        resume_entries=[ResumeSourceEntry(
            entry_id="source-entry-1", kind="other", heading="用户确认的补充经历",
            evidence_ids=["supplement_1"], original_bullets=[answer],
        )],
    )
    public = _public_profile(profile, CandidateContact(profile_id="profile_1"))
    assert public.resume_entries[0].kind == "project"
    assert public.resume_entries[0].heading == "股票信息数据采集程序"
    assert supplement_entry_metadata("我有过爬虫经验，写过股票信息数据爬虫，项目仓库地址：https://example.com") == (
        "project", "股票信息数据爬虫",
    )


def test_resume_parser_keeps_project_title_date_and_bullets_together():
    profile = build_profile("""王科
教育经历
湖南人文科技学院
2025-06 - 2027-09
计算机科学与技术 | 本科
项目经历
Text-to-SQL 智能Agent系统
2026-02 - 2026-03
项目地址：https://gitee.com/wk132/data-agent.git
1. 设计并实现基于 LangGraph 的多节点 Agent 管线。
RAG 知识库系统
2026-04 - 2027-05
项目地址：https://gitee.com/wk132/rag-knowledge-base.git
1. 设计并实现基于 LangGraph 的 RAG 知识库系统。
""")
    projects = [entry for entry in profile.resume_entries if entry.kind == "project"]
    assert [(entry.heading, entry.date_range) for entry in projects] == [
        ("Text-to-SQL 智能Agent系统", "2026-02 - 2026-03"),
        ("RAG 知识库系统", "2026-04 - 2027-05"),
    ]
    assert len(projects[0].evidence_ids) == 4
    education = next(entry for entry in profile.resume_entries if entry.kind == "education")
    assert education.heading == "湖南人文科技学院"
    assert education.date_range == "2025-06 - 2027-09"


def test_old_draft_metadata_is_repaired_from_source_entries():
    profile = build_profile("""王科
教育经历
湖南人文科技学院
2025-06 - 2027-09
计算机科学与技术 | 本科
项目经历
Text-to-SQL 智能Agent系统
2026-02 - 2026-03
项目地址：https://gitee.com/wk132/data-agent.git
1. 设计并实现基于 LangGraph 的多节点 Agent 管线。
""")
    target = TargetJob(
        target_job_id="target_1", profile_id=profile.profile_id, source_type="pasted",
        company="星图科技", title="Python 后端工程师", cleaned_text="Python 后端工程师",
        created_at=datetime.now(UTC).isoformat(),
    )
    version = _version(target)
    version.sections = [
        ResumeSection(
            section_id="section_project", title="项目经历", entries=[ResumeEntry(
                entry_id="entry_old", heading="2026-02 - 2026-03",
                date_range="2026-02 - 2026-03",
                evidence_ids=[item.block_id for item in profile.evidence if "项目" in item.block_id][:4],
                bullets=[ResumeBullet(
                    bullet_id="bullet_1", text="设计并实现基于 LangGraph 的多节点 Agent 管线。",
                    evidence_ids=[item.block_id for item in profile.evidence if "多节点" in item.quote],
                    provenance="uploaded_resume",
                )],
                )],
        )
    ]
    assert _repair_draft_source_metadata(version, profile) is True
    entry = version.sections[0].entries[0]
    assert entry.heading == "Text-to-SQL 智能Agent系统"
    assert entry.date_range == "2026-02 - 2026-03"
    assert entry.links == ["https://gitee.com/wk132/data-agent.git"]


def test_export_qa_reports_page_count_and_ats_state(tmp_path: Path):
    target = TargetJob(
        target_job_id="target_1", profile_id="profile_1", source_type="snapshot",
        company="星图科技", title="Python 后端工程师", cleaned_text="Python 后端工程师",
        created_at=datetime.now(UTC).isoformat(),
    )
    version = _version(target)
    docx_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("张三")
    document.add_paragraph("13800138000 · zhangsan@example.com")
    document.add_paragraph("星图科技 · Python 后端工程师")
    document.save(docx_path)
    pdf_path = tmp_path / "resume.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(pdf_path)

    report = inspect_export(version, target, "technical", docx_path, pdf_path)
    assert report.page_count == 1
    assert report.ats_extractable is False
    assert report.status == "WARN"
    assert any(item.id == "pdf_text_layer" and item.status == "WARN" for item in report.checks)


def test_export_qa_blocks_more_than_two_pages(tmp_path: Path):
    target = TargetJob(
        target_job_id="target_1", profile_id="profile_1", source_type="snapshot",
        company="星图科技", title="Python 后端工程师", cleaned_text="Python 后端工程师",
        created_at=datetime.now(UTC).isoformat(),
    )
    version = _version(target)
    docx_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("张三")
    document.add_paragraph("星图科技 · Python 后端工程师")
    document.save(docx_path)
    pdf_path = tmp_path / "resume.pdf"
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=595, height=842)
    with pdf_path.open("wb") as stream:
        writer.write(stream)

    report = inspect_export(version, target, "business", docx_path, pdf_path)
    assert report.status == "FAIL"
    assert any(item.id == "page_count" and item.status == "FAIL" for item in report.checks)
