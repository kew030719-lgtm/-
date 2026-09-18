from __future__ import annotations

import hashlib
import html
import json
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from .agent import AgentService, _extract_json
from .database import Database, now_iso
from .resume import (
    EMAIL_RE, PHONE_RE, SKILLS, ResumeError, ResumeGenerationError, extract_candidate_contact,
    sanitize_candidate_contact,
    skill_is_grounded,
)
from .schemas import (
    CandidateContact, CandidateProfile, Evidence, JobSnapshot, ResumeBullet,
    ResumeDraftVersion, ResumeEntry, ResumeExport, ResumeSection, ResumeSourceEntry, ResumeTailoring,
    SupplementalEvidence, TailoringQuestion, TargetJob,
)
from .sites.base import is_benefit_label


def _stable_id(prefix: str, value: str) -> str:
    return f"{prefix}_{hashlib.sha256(value.encode()).hexdigest()[:12]}"


def _job_blocks(target_id: str, values: list[str]) -> list[Evidence]:
    blocks = []
    for index, value in enumerate(values, start=1):
        value = value.strip()
        if value:
            blocks.append(Evidence(
                source_type="job", source_id=target_id,
                block_id=f"target-{index:02d}-{hashlib.sha256(value.encode()).hexdigest()[:10]}",
                quote=value, section="目标岗位", provenance="job_snapshot",
            ))
    return blocks


def target_from_snapshot(profile_id: str, snapshot: JobSnapshot,
                         source_type: str = "snapshot", target_id: str | None = None) -> TargetJob:
    target_id = target_id or f"target_{uuid4().hex[:12]}"
    required_skills = [item for item in snapshot.required_skills if not is_benefit_label(item)]
    values = [snapshot.title, snapshot.company, *snapshot.responsibilities, *required_skills]
    return TargetJob(
        target_job_id=target_id, profile_id=profile_id, source_type=source_type,
        # A target derived from a stored snapshot inherits that snapshot's site, so
        # a 智联 posting is never later described as coming from BOSS.
        site=snapshot.site or "boss",
        snapshot_id=snapshot.snapshot_id, company=snapshot.company, title=snapshot.title,
        url=snapshot.canonical_url, cleaned_text=snapshot.cleaned_text,
        responsibilities=snapshot.responsibilities, required_skills=required_skills,
        blocks=_job_blocks(target_id, values), created_at=now_iso(),
    )


def target_from_pasted(profile_id: str, company: str, title: str, content: str,
                       url: str = "") -> TargetJob:
    cleaned = re.sub(r"\s+", " ", content).strip()
    if len(cleaned) < 30:
        raise ResumeError("岗位描述过短，请粘贴职责、技能和任职要求")
    target_id = f"target_{uuid4().hex[:12]}"
    lowered = cleaned.lower()
    required = sorted({display for token, display in SKILLS.items() if token in lowered})
    lines = [line.strip() for line in re.split(r"[\n；;。]", content) if line.strip()]
    responsibilities = lines[:30]
    return TargetJob(
        target_job_id=target_id, profile_id=profile_id, source_type="pasted",
        company=company.strip(), title=title.strip(), url=url.strip(), cleaned_text=cleaned[:20_000],
        responsibilities=responsibilities, required_skills=required,
        blocks=_job_blocks(target_id, [title, company, *lines]), created_at=now_iso(),
    )


def _evidence_score(evidence: Evidence, target: TargetJob) -> int:
    text = evidence.quote.lower()
    terms = [*target.required_skills, *re.findall(r"[A-Za-z][A-Za-z0-9+#.]{1,20}", target.cleaned_text)]
    return sum(12 for term in terms if term.lower() in text) + min(len(text), 80) // 20


def _is_resume_noise(evidence: Evidence) -> bool:
    text = evidence.quote.strip(" ：:").lower()
    exact = {
        "基本信息", "个人信息", "个人简历", "求职简历", "简历", "联系方式",
        "教育", "教育经历", "技能", "专业技能", "技术栈", "项目", "项目经历",
        "经历", "工作经历", "实习经历",
    }
    noisy_phrases = ("温馨提示", "简历模板", "虚构示例", "请根据实际情况", "仅供参考")
    return text in exact or any(phrase in text for phrase in noisy_phrases)


def _questions(profile: CandidateProfile, target: TargetJob,
               supplements: list[SupplementalEvidence]) -> tuple[list[TailoringQuestion], list[str]]:
    known = " ".join([*(item.quote for item in profile.evidence), *(item.quote for item in supplements)]).lower()
    required_skills = [skill for skill in target.required_skills if not is_benefit_label(skill)]
    missing = [skill for skill in required_skills if skill.lower() not in known]
    questions = [TailoringQuestion(
        question_id=_stable_id("question", f"{target.target_job_id}:{skill}"),
        question=f"岗位要求 {skill}。你是否在真实项目中使用过？如果有，请说明使用场景、你的职责和结果；没有可以跳过。",
        requirement=skill,
    ) for skill in missing[:5]]
    if len(questions) < 5:
        signals = ("年以上", "至少", "经验", "具备", "熟悉", "能够", "负责", "设计", "建立", "推动", "分析", "评测")
        boilerplate = ("在这里", "真诚地邀请", "技术氛围", "职业生涯", "平台是", "一起驱动", "共同探索", "站在")
        concepts = (
            "产品管理", "项目管理", "用户研究", "数据分析", "问题归因", "评测体系",
            "评测数据集", "评分标准", "回归测试", "自动化评测", "人工评测", "A/B 实验",
            "大模型", "工具调用", "跨团队协作", "产品路线图", "模型选型",
        )
        candidates = []
        for index, raw in enumerate(target.responsibilities):
            requirement = re.sub(r"^\s*\d+[、.．)）]\s*", "", raw).strip()
            if not 12 <= len(requirement) <= 220 or any(value in requirement for value in boilerplate):
                continue
            score = sum(3 if value in {"年以上", "至少", "经验", "具备", "熟悉", "能够"} else 1
                        for value in signals if value in requirement)
            if score == 0:
                continue
            lowered = requirement.lower()
            terms = {display.lower() for token, display in SKILLS.items() if token in lowered}
            terms.update(token.lower() for token in re.findall(r"[A-Za-z][A-Za-z0-9+#./-]{1,24}", requirement))
            terms.update(value.lower() for value in concepts if value.lower() in lowered)
            covered = sum(1 for term in terms if term in known)
            if terms and covered >= max(2, round(len(terms) * 0.6)):
                continue
            candidates.append((-score, index, requirement))
        for _score, _index, requirement in sorted(candidates):
            if len(questions) >= 5:
                break
            if any(skill.lower() in requirement.lower() for skill in missing):
                continue
            questions.append(TailoringQuestion(
                question_id=_stable_id("question", f"{target.target_job_id}:{requirement}"),
                question=(
                    f"岗位强调“{requirement[:90]}”。你的真实经历中是否有能证明这一点的案例？"
                    "如有请说明使用场景、个人贡献、交付结果和可证实数据；没有可以跳过。"
                ),
                requirement=requirement,
            ))
    return questions, missing


def _as_bullet(evidence: Evidence, target: TargetJob, index: int,
               provenance: str | None = None) -> ResumeBullet:
    return ResumeBullet(
        bullet_id=f"bullet_{uuid4().hex[:10]}", text=evidence.quote,
        evidence_ids=[evidence.block_id],
        provenance=provenance or evidence.provenance or "uploaded_resume",
        priority=min(100, 40 + _evidence_score(evidence, target) + max(0, 12 - index)),
    )


def _fallback_draft(profile: CandidateProfile, contact: CandidateContact, target: TargetJob,
                    tailoring: ResumeTailoring, draft_id: str) -> ResumeDraftVersion:
    usable = [item for item in profile.evidence if not _is_resume_noise(item)]
    ordered = sorted(usable, key=lambda item: _evidence_score(item, target), reverse=True)
    by_section: dict[str, list[Evidence]] = {}
    for item in ordered:
        by_section.setdefault(item.section or "其他", []).append(item)

    summary_pool = [item for item in ordered if item.section in {"经历", "项目", "用户补充", "概览"}]
    summary_sources = summary_pool[:2] or ordered[:2]
    summary = [_as_bullet(item, target, index) for index, item in enumerate(summary_sources)]
    skills = []
    for index, skill in enumerate(profile.skills):
        source = next((item for item in profile.evidence if skill.lower() in item.quote.lower()), None)
        if source:
            skills.append(ResumeBullet(
                bullet_id=f"skill_{uuid4().hex[:10]}", text=skill,
                evidence_ids=[source.block_id], provenance=source.provenance or "uploaded_resume",
                priority=90 if skill in target.required_skills else 55 - min(index, 20),
            ))
    skills = sorted(skills, key=lambda item: item.priority, reverse=True)[:8]
    sections = []
    labels = (
        ("经历", "工作经历", 5), ("项目", "项目经历", 4),
        ("用户补充", "补充经历", 3), ("教育", "教育经历", 2),
        ("技能", "补充技能", 2), ("概览", "其他信息", 1),
    )
    used = {item.block_id for item in summary_sources}
    for key, title, limit in labels:
        items = [item for item in by_section.get(key, []) if item.block_id not in used]
        if not items:
            continue
        bullets = [_as_bullet(item, target, index) for index, item in enumerate(items[:limit])]
        sections.append(ResumeSection(section_id=f"section_{key}", title=title, bullets=bullets))
    return ResumeDraftVersion(
        version_id=f"version_{uuid4().hex[:12]}", draft_id=draft_id, version=1,
        tailoring_id=tailoring.tailoring_id, profile_id=profile.profile_id,
        target_job_id=target.target_job_id, headline=target.title,
        summary=summary, skills=skills,
        sections=sections, contact=contact, source="fallback", created_at=now_iso(),
        change_log=[{"kind": "generated", "target": f"{target.company} · {target.title}"}],
    )


def _all_bullets(version: ResumeDraftVersion) -> list[ResumeBullet]:
    values = [*version.summary, *version.skills]
    for section in version.sections:
        values.extend(section.bullets)
        for entry in section.entries:
            values.extend(entry.bullets)
    return values


def _public_profile(profile: CandidateProfile, contact: CandidateContact) -> CandidateProfile:
    value = profile.model_copy(deep=True)
    private = {item for item in (contact.name, contact.phone, contact.email, contact.location) if item}
    value.evidence = [item for item in value.evidence if not _is_resume_noise(item) and not (
        PHONE_RE.search(item.quote) or EMAIL_RE.search(item.quote) or
        any(secret in item.quote for secret in private)
    )]
    safe = {item.block_id: item.quote for item in value.evidence}
    entries = []
    for entry in value.resume_entries:
        entry.evidence_ids = [item for item in entry.evidence_ids if item in safe]
        entry.original_bullets = [safe[item] for item in entry.evidence_ids]
        if entry.evidence_ids:
            entries.append(entry)
    value.resume_entries = entries
    return value


def validate_draft(version: ResumeDraftVersion, profile: CandidateProfile) -> None:
    sources = {item.block_id: item for item in profile.evidence}

    def validate_text(text: str, evidence_ids: list[str], label: str) -> None:
        if not evidence_ids or any(evidence_id not in sources for evidence_id in evidence_ids):
            raise ResumeError(f"简历内容引用了无效证据：{label}")
        quotes = " ".join(sources[item].quote for item in evidence_ids)
        for number in re.findall(r"\d+(?:\.\d+)?%?", text):
            if number not in quotes:
                raise ResumeError(f"简历内容出现未经证实的数字：{number}")
        for entity in re.findall(r"[A-Za-z0-9\u4e00-\u9fff]{2,30}(?:有限责任公司|有限公司|大学|学院|研究院)", text):
            if entity not in quotes:
                raise ResumeError(f"简历内容出现未经证实的机构：{entity}")
        for token, display in SKILLS.items():
            if token in text.lower() and not skill_is_grounded(token, quotes):
                raise ResumeError(f"简历内容出现未经证实的技能：{display}")

    for bullet in _all_bullets(version):
        validate_text(bullet.text, bullet.evidence_ids, bullet.bullet_id)
        if any(phrase in bullet.text.lower() for phrase in ("温馨提示", "简历模板", "虚构示例", "请根据实际情况", "仅供参考")):
            raise ResumeError("简历内容包含模板提示语")
    for section in version.sections:
        for entry in section.entries:
            validate_text(" ".join(filter(None, [entry.heading, entry.subheading, entry.date_range])),
                          entry.evidence_ids, entry.entry_id)
            entry_quotes = " ".join(sources[item].quote for item in entry.evidence_ids)
            for metadata in (entry.heading, entry.subheading, entry.date_range):
                if metadata and metadata not in entry_quotes:
                    raise ResumeError(f"简历条目出现未经证实的标题或日期：{metadata}")
            allowed = set(entry.evidence_ids)
            if any(not set(bullet.evidence_ids).issubset(allowed) for bullet in entry.bullets):
                raise ResumeError(f"简历条目混入了其他经历的事实：{entry.heading}")


def assess_draft_quality(version: ResumeDraftVersion, target: TargetJob) -> None:
    bullets = _all_bullets(version)
    normalized = [re.sub(r"\W+", "", item.text).lower() for item in bullets]
    duplicates = len(normalized) - len(set(normalized))
    text = " ".join(item.text for item in bullets)
    characters = len(text) + sum(len(entry.heading + entry.subheading + entry.date_range)
                                 for section in version.sections for entry in section.entries)
    report = version.quality_report
    report.evidence_coverage = 100 if bullets and all(item.evidence_ids for item in bullets) else 0
    report.duplicate_count = duplicates
    report.estimated_pages = round(max(0.5, characters / 1600), 1)
    report.uncovered_requirements = [skill for skill in target.required_skills if skill.lower() not in text.lower()]
    issues = list(dict.fromkeys(report.issues))
    if duplicates:
        issues.append(f"存在 {duplicates} 条重复内容")
    if report.estimated_pages > 1.2:
        issues.append("内容超过默认一页篇幅")
    if not version.sections or not bullets:
        issues.append("简历缺少可用的经历内容")
    report.issues = list(dict.fromkeys(issues))
    report.passed = bool(
        report.passed and report.evidence_coverage == 100 and not duplicates
        and report.estimated_pages <= 1.2 and version.sections and bullets
        and min(report.relevance_score, report.specificity_score,
                report.structure_score, report.conciseness_score) >= 70
    )
    if not report.passed:
        raise ResumeError("定向简历未通过质量检查：" + "；".join(report.issues[:5] or ["模型评分未达到 70 分"]))


class TailoringService:
    def __init__(self, database: Database, agent: AgentService, data_dir: Path):
        self.database = database
        self.agent = agent
        self.export_dir = data_dir / "exports"

    def create_tailoring(self, profile_id: str, target_job_id: str,
                         conversation_id: str | None = None) -> ResumeTailoring:
        profile = self.database.get_profile(profile_id)
        target = self.database.get_target_job(target_job_id)
        if not profile or not target or target.profile_id != profile_id:
            raise ResumeError("目标岗位与候选人画像不匹配")
        questions, missing = _questions(profile, target, self.database.list_supplemental_evidence(profile_id))
        stamp = now_iso()
        tailoring = ResumeTailoring(
            tailoring_id=f"tailor_{uuid4().hex[:12]}", conversation_id=conversation_id,
            profile_id=profile_id, target_job_id=target_job_id,
            status="COLLECTING" if questions else "READY", questions=questions,
            missing_requirements=missing, created_at=stamp, updated_at=stamp,
        )
        return self.database.save_tailoring(tailoring)

    def restore_missing_questions(self, tailoring: ResumeTailoring) -> ResumeTailoring:
        """Backfill questions for tasks created before JD requirement fallback existed."""
        has_benefit_questions = any(is_benefit_label(item.requirement) for item in tailoring.questions)
        if (tailoring.questions and not has_benefit_questions) or tailoring.status == "SUCCEEDED":
            return tailoring
        profile = self.database.get_profile(tailoring.profile_id)
        target = self.database.get_target_job(tailoring.target_job_id)
        if not profile or not target:
            return tailoring
        questions, missing = _questions(
            profile, target, self.database.list_supplemental_evidence(profile.profile_id),
        )
        if not questions and not has_benefit_questions:
            return tailoring
        previous = {item.question_id: item for item in tailoring.questions}
        for question in questions:
            if question.question_id in previous:
                question.status = previous[question.question_id].status
                question.answer = previous[question.question_id].answer
        tailoring.questions = questions
        tailoring.missing_requirements = missing
        tailoring.status = "READY" if all(item.status != "PENDING" for item in questions) else "COLLECTING"
        tailoring.task_id = None
        tailoring.error = None
        return self.database.save_tailoring(tailoring)

    def save_answers(self, tailoring_id: str, answers: dict[str, str | None]) -> ResumeTailoring:
        tailoring = self.database.get_tailoring(tailoring_id)
        if not tailoring:
            raise ResumeError("定向简历任务不存在")
        if tailoring.status not in {"COLLECTING", "READY", "FAILED", "FAILED_VALIDATION"}:
            raise ResumeError("当前状态不能修改追问答案")
        if tailoring.status in {"FAILED", "FAILED_VALIDATION"}:
            tailoring.task_id = None
            tailoring.error = None
        valid_ids = {question.question_id for question in tailoring.questions}
        if set(answers) - valid_ids:
            raise ResumeError("包含不属于当前任务的追问")
        for question in tailoring.questions:
            if question.question_id not in answers:
                continue
            answer = (answers[question.question_id] or "").strip()
            if len(answer) > 1000:
                raise ResumeError("单条补充说明不能超过 1000 字")
            question.answer = answer or None
            question.status = "ANSWERED" if answer else "SKIPPED"
        tailoring.status = "READY" if all(item.status != "PENDING" for item in tailoring.questions) else "COLLECTING"
        return self.database.save_tailoring(tailoring)

    def confirm_facts(self, tailoring: ResumeTailoring) -> CandidateProfile:
        profile = self.database.get_profile(tailoring.profile_id)
        if not profile:
            raise ResumeError("候选人画像不存在")
        stamp = now_iso()
        existing = {item.block_id for item in profile.evidence}
        for question in tailoring.questions:
            if question.status != "ANSWERED" or not question.answer:
                continue
            evidence_id = _stable_id("resume-supplement", f"{profile.profile_id}:{question.answer}")
            stored = SupplementalEvidence(
                evidence_id=evidence_id, profile_id=profile.profile_id,
                tailoring_id=tailoring.tailoring_id, question_id=question.question_id,
                quote=question.answer, created_at=stamp,
            )
            self.database.save_supplemental_evidence(stored)
            if evidence_id not in existing:
                profile.evidence.append(Evidence(
                    source_type="resume", source_id=profile.profile_id, block_id=evidence_id,
                    quote=question.answer, section="用户补充", provenance="user_confirmed",
                ))
                profile.resume_entries.append(ResumeSourceEntry(
                    entry_id=f"source-entry-{evidence_id}", kind="other",
                    heading="用户确认的补充经历", evidence_ids=[evidence_id],
                    original_bullets=[question.answer],
                ))
                existing.add(evidence_id)
        self.database.save_profile(profile)
        return profile

    async def generate(self, tailoring_id: str) -> ResumeDraftVersion:
        tailoring = self.database.get_tailoring(tailoring_id)
        if not tailoring or tailoring.status not in {"READY", "GENERATING"}:
            raise ResumeError("追问尚未完成或任务状态无效")
        profile = self.confirm_facts(tailoring)
        target = self.database.get_target_job(tailoring.target_job_id)
        if not target:
            raise ResumeError("目标岗位不存在")
        contact = sanitize_candidate_contact(self.database.get_contact(profile.profile_id))
        if not any((contact.name, contact.phone, contact.email, contact.location)):
            candidate_copy = profile.model_copy(deep=True)
            contact = sanitize_candidate_contact(extract_candidate_contact(candidate_copy))
        self.database.save_contact(contact)
        public_profile = _public_profile(profile, contact)
        draft_id = tailoring.draft_id or f"draft_{uuid4().hex[:12]}"
        version = _fallback_draft(public_profile, contact, target, tailoring, draft_id)
        latest = self.database.get_latest_draft(draft_id)
        if latest:
            version.version = latest.version + 1
        if not self.agent.settings.api_key:
            tailoring.status = "FAILED"
            tailoring.error = "尚未配置可用的大模型，已停止生成，未展示低质量整理版"
            self.database.save_tailoring(tailoring)
            raise ResumeGenerationError(tailoring.error)
        try:
            version = await self.agent.tailor_resume(
                public_profile, target, tailoring, contact, version,
                validator=lambda draft: validate_draft(draft, profile),
            )
        except Exception as exc:
            tailoring.status = "FAILED"
            reason = str(exc).strip() or ("请求超时" if isinstance(exc, TimeoutError) else type(exc).__name__)
            tailoring.error = f"大模型生成或质量返工失败：{reason}"
            self.database.save_tailoring(tailoring)
            raise ResumeGenerationError(tailoring.error) from exc
        try:
            validate_draft(version, profile)
            assess_draft_quality(version, target)
        except ResumeError as exc:
            version.validation_status = "FAILED_VALIDATION"
            self.database.save_draft_version(version)
            tailoring.status = "FAILED_VALIDATION"
            tailoring.error = str(exc)
            self.database.save_tailoring(tailoring)
            raise
        self.database.save_draft_version(version)
        tailoring.status = "SUCCEEDED"
        tailoring.error = None
        tailoring.draft_id = draft_id
        self.database.save_tailoring(tailoring)
        return version

    def save_user_version(self, draft_id: str, value: ResumeDraftVersion) -> ResumeDraftVersion:
        previous = self.database.get_latest_draft(draft_id)
        if not previous or value.profile_id != previous.profile_id or value.target_job_id != previous.target_job_id:
            raise ResumeError("草稿版本上下文不匹配")
        profile = self.database.get_profile(previous.profile_id)
        if not profile:
            raise ResumeError("候选人画像不存在")
        previous_text = {item.bullet_id: item.text for item in _all_bullets(previous)}
        change_log = []
        if value.headline != previous.headline:
            change_log.append({"kind": "headline", "before": previous.headline, "after": value.headline})
        value.contact.profile_id = previous.profile_id
        for field in ("name", "phone", "email", "location"):
            before = getattr(previous.contact, field)
            after = getattr(value.contact, field)
            if before != after:
                change_log.append({
                    "kind": "contact", "field": field,
                    "before": "已填写" if before else "未填写",
                    "after": "已填写" if after else "未填写",
                })
        for bullet in _all_bullets(value):
            old = previous_text.get(bullet.bullet_id)
            if old == bullet.text:
                continue
            evidence_id = _stable_id("resume-edit", f"{profile.profile_id}:{bullet.text}")
            if not any(item.block_id == evidence_id for item in profile.evidence):
                profile.evidence.append(Evidence(
                    source_type="resume", source_id=profile.profile_id, block_id=evidence_id,
                    quote=bullet.text, section="用户编辑", provenance="user_edited",
                ))
            bullet.evidence_ids = [evidence_id]
            bullet.provenance = "user_edited"
            change_log.append({"bullet_id": bullet.bullet_id, "before": old or "", "after": bullet.text})
        value.version_id = f"version_{uuid4().hex[:12]}"
        value.version = previous.version + 1
        value.draft_id = draft_id
        value.tailoring_id = previous.tailoring_id
        value.source = "user_edit"
        value.change_log = change_log
        value.created_at = now_iso()
        value.contact = sanitize_candidate_contact(value.contact)
        self.database.save_contact(value.contact)
        self.database.save_profile(profile)
        validate_draft(value, profile)
        return self.database.save_draft_version(value)

    def render_docx(self, version: ResumeDraftVersion, target: TargetJob,
                    template: str, output: Path) -> None:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        styles = document.styles
        styles["Normal"].font.name = "Noto Sans CJK SC"
        styles["Normal"].font.size = Pt(9.5 if template == "business" else 10)
        styles["Normal"].paragraph_format.space_after = Pt(0)
        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(1)
        run = title.add_run(version.contact.name or "个人简历")
        run.font.name = "Noto Sans CJK SC"
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        contacts = "  ·  ".join(filter(None, [version.contact.phone, version.contact.email, version.contact.location]))
        if contacts:
            paragraph = document.add_paragraph(contacts)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(1)
        subtitle = document.add_paragraph(f"目标岗位：{target.company} · {version.headline}")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(1)

        def add_heading(text: str) -> None:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0, 0, 0)

        def add_bullets(values: list[ResumeBullet]) -> None:
            for item in sorted(values, key=lambda bullet: bullet.priority, reverse=True):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(0.5)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.add_run(item.text)

        def add_skills(values: list[ResumeBullet]) -> None:
            paragraph = document.add_paragraph(" · ".join(
                item.text for item in sorted(values, key=lambda bullet: bullet.priority, reverse=True)
            ))
            paragraph.paragraph_format.space_after = Pt(1)

        priorities = {"项目经历": 0, "工作经历": 1, "补充技能": 2, "教育经历": 3, "其他信息": 4} if template == "technical" else {"工作经历": 0, "项目经历": 1, "教育经历": 2, "补充技能": 3, "其他信息": 4}
        ordered = sorted(version.sections, key=lambda item: priorities.get(item.title, 9))
        if template == "technical":
            add_heading("专业技能")
            add_skills(version.skills)
            add_heading("个人概述")
            add_bullets(version.summary)
        else:
            add_heading("个人概述")
            add_bullets(version.summary)
        for item in ordered:
            add_heading(item.title)
            add_bullets(item.bullets)
            for entry in item.entries:
                paragraph = document.add_paragraph()
                paragraph.add_run(entry.heading).bold = True
                if entry.subheading:
                    paragraph.add_run(f"  {entry.subheading}")
                if entry.date_range:
                    paragraph.add_run(f"  {entry.date_range}")
                add_bullets(entry.bullets)
        if template == "business" and version.skills:
            add_heading("专业技能")
            add_skills(version.skills)
        document.core_properties.title = f"{target.company} {version.headline} 定向简历"
        document.core_properties.author = version.contact.name
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)

    @staticmethod
    def render_html(version: ResumeDraftVersion, target: TargetJob, template: str) -> str:
        def bullets(values: list[ResumeBullet]) -> str:
            items = "".join(f"<li>{html.escape(item.text)}</li>" for item in sorted(values, key=lambda value: value.priority, reverse=True))
            return f"<ul>{items}</ul>" if items else ""
        def skills(values: list[ResumeBullet]) -> str:
            value = " · ".join(html.escape(item.text) for item in sorted(values, key=lambda item: item.priority, reverse=True))
            return f"<p class='skills'>{value}</p>" if value else ""
        sections = []
        if template == "technical":
            sections.extend([("专业技能", skills(version.skills)), ("个人概述", bullets(version.summary))])
        else:
            sections.append(("个人概述", bullets(version.summary)))
        priorities = {"项目经历": 0, "工作经历": 1, "补充技能": 2, "教育经历": 3, "其他信息": 4} if template == "technical" else {"工作经历": 0, "项目经历": 1, "教育经历": 2, "补充技能": 3, "其他信息": 4}
        for section in sorted(version.sections, key=lambda item: priorities.get(item.title, 9)):
            body = bullets(section.bullets)
            for entry in section.entries:
                body += f"<h3>{html.escape(entry.heading)} <small>{html.escape(entry.subheading)} {html.escape(entry.date_range)}</small></h3>{bullets(entry.bullets)}"
            sections.append((section.title, body))
        if template == "business":
            sections.append(("专业技能", skills(version.skills)))
        body = "".join(f"<section><h2>{html.escape(title)}</h2>{content}</section>" for title, content in sections if content)
        contact = " · ".join(filter(None, [version.contact.phone, version.contact.email, version.contact.location]))
        accent = "#102dff" if template == "technical" else "#193b2d"
        return f"""<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><style>
        @page{{size:A4;margin:13mm 16mm}}*{{box-sizing:border-box}}body{{font-family:'Noto Sans CJK SC','Microsoft YaHei',sans-serif;color:#101828;font-size:10pt;line-height:1.4;margin:0}}header{{text-align:center;margin-bottom:4mm}}h1{{font-size:22pt;margin:0;color:#000}}header p{{margin:1.5mm 0;color:#475467}}.target{{color:{accent};font-weight:700}}h2{{font-size:12pt;color:#000;border-bottom:1.5px solid {accent};padding-bottom:1mm;margin:3mm 0 1.2mm}}h3{{font-size:10pt;margin:2mm 0 1mm}}h3 small{{font-weight:400;color:#667085}}ul{{margin:0;padding-left:5mm}}li{{margin:0 0 .8mm;break-inside:avoid}}.skills{{margin:0;line-height:1.6}}section{{break-inside:auto}}
        </style></head><body><header><h1>{html.escape(version.contact.name or '个人简历')}</h1><p>{html.escape(contact)}</p><p class='target'>目标岗位：{html.escape(target.company)} · {html.escape(version.headline)}</p></header>{body}</body></html>"""

    async def render_pdf(self, html_value: str, output: Path) -> None:
        from playwright.async_api import async_playwright
        output.parent.mkdir(parents=True, exist_ok=True)
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            page = await browser.new_page()
            await page.set_content(html_value, wait_until="load")
            await page.pdf(path=str(output), format="A4", print_background=True,
                           margin={"top": "0", "right": "0", "bottom": "0", "left": "0"})
            await browser.close()

    async def export(self, export_id: str) -> ResumeExport:
        export = self.database.get_resume_export(export_id)
        if not export:
            raise ResumeError("导出任务不存在")
        version = self.database.get_draft_version(export.version_id)
        if not version or version.validation_status != "VALID":
            raise ResumeError("简历版本不存在或未通过校验")
        target = self.database.get_target_job(version.target_job_id)
        if not target:
            raise ResumeError("目标岗位不存在")
        safe = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", f"{target.company}_{target.title}").strip("_")[:80]
        folder = self.export_dir / export.export_id
        docx_path = folder / f"{safe}_简历_v{version.version}_{export.template}.docx"
        pdf_path = folder / f"{safe}_简历_v{version.version}_{export.template}.pdf"
        self.render_docx(version, target, export.template, docx_path)
        await self.render_pdf(self.render_html(version, target, export.template), pdf_path)
        export.status = "SUCCEEDED"
        export.docx_path = str(docx_path)
        export.pdf_path = str(pdf_path)
        return self.database.save_resume_export(export)
